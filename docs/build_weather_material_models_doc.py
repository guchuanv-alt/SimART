#!/usr/bin/env python3
"""Build the concise two-function weather material model note."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "SimART_天气材质两参数简化模型.docx"


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def table(doc: Document, headers: list[str], rows: list[tuple[str, ...]]) -> None:
    result = doc.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    for index, title in enumerate(headers):
        cell = result.rows[0].cells[index]
        cell.text = title
        shade(cell, "D9EAF7")
        cell.paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = result.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    doc.add_paragraph()


def equation(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Cambria Math"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
    run.font.size = Pt(10.5)


def note(doc: Document, title: str, text: str) -> None:
    result = doc.add_table(rows=1, cols=1)
    result.style = "Table Grid"
    cell = result.cell(0, 0)
    shade(cell, "FFF2CC")
    p = cell.paragraphs[0]
    p.add_run(title + "：").bold = True
    p.add_run(text)
    doc.add_paragraph()


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)

    title = doc.add_heading("SimART 天气材质两参数简化模型", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("用途：由 Agent 根据天气描述，为当前 BigCity XML 的每个材质槽计算相对介电常数 εr 和电导率 σ。固定仿真频率：3.5 GHz。")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    note(doc, "这份文档只保留工程第一版", "只考虑材质类别、材料内部湿润度、地面积水厚度、屋顶水膜厚度、积雪厚度、结冰厚度、温度和表面角色。不会引入降雨历时、渗透率、散射或极化等额外变量。")

    doc.add_heading("1. Agent 最终要做什么", level=1)
    doc.add_paragraph("用户说天气；DeepSeek 只把自然语言整理为输入变量；本地程序调用下面两个函数；程序写出一份新的 XML 并让 GUI 选择该 XML。大模型不直接编造 εr 或 σ。")
    equation(doc, "εr_effective = f(material, u, dw, dr, ds, di, T, role)")
    equation(doc, "σ_effective  = g(material, u, dw, dr, ds, di, T, role)")

    doc.add_heading("2. 所有输入值是什么意思", level=1)
    table(doc, ["变量", "含义", "单位/范围", "由谁给出"], [
        ("material", "材质类别，例如混凝土、砖、玻璃、金属、木材、道路沥青、屋顶沥青、草地。", "文本类别", "XML 槽名、贴图名和对象名自动识别"),
        ("u", "材料内部湿润度。0 表示该材料干燥；1 表示该材料达到预先定义的充分湿润状态。", "0 到 1", "Agent 从“晴天/小雨/大雨/很潮湿”等描述推断，后续可由传感器直接提供"),
        ("dw", "水平地面的积水或水膜厚度。", "mm", "用户或天气/传感器输入"),
        ("dr", "屋顶水膜或屋顶积水厚度。", "mm", "用户或天气/传感器输入"),
        ("ds", "积雪厚度。", "mm", "用户或天气/传感器输入"),
        ("di", "冰层厚度。", "mm", "用户或天气/传感器输入"),
        ("T", "表面温度；决定水、冰、雪使用哪组介电参数。", "摄氏度", "用户或气象数据"),
        ("role", "表面角色：ground、roof 或 wall。", "类别", "由槽名/对象名识别"),
    ])

    doc.add_heading("3. 第一步：材料本体因内部湿润而变化", level=1)
    doc.add_paragraph("每一种材料都在材料表中保存四个端点值：干燥状态的 εr_dry、σ_dry，以及充分湿润状态的 εr_wet、σ_wet。Agent 不让大模型临时决定这些端点。")
    equation(doc, "εr_bulk = εr_dry(material) + u × [εr_wet(material) − εr_dry(material)]")
    equation(doc, "σ_bulk  = σ_dry(material)  + u × [σ_wet(material) − σ_dry(material)]")
    table(doc, ["符号", "意思"], [
        ("εr_dry", "该材料在干燥/典型状态下的相对介电常数。"),
        ("σ_dry", "该材料在干燥/典型状态下的电导率，单位 S/m。"),
        ("εr_wet", "该材料本体充分湿润后的相对介电常数；不包含额外积水或冰雪。"),
        ("σ_wet", "该材料本体充分湿润后的电导率；不包含额外积水或冰雪。"),
        ("u", "内部湿润度，因此本体参数始终在干端点与湿端点之间。"),
    ])
    doc.add_paragraph("当前端点表与程序完全一致，固定在 3.5 GHz、20°C。温度在第一版只改变水、冰、雪覆盖层的参数；不单独改变材料本体端点。")
    table(doc, ["材料类别", "εr_dry", "σ_dry (S/m)", "εr_wet", "σ_wet (S/m)", "说明"], [
        ("混凝土", "5.2400", "0.123087", "10.3837", "0.300201", "干端点来自 ITU 混凝土；湿端点为本地工程端点。"),
        ("砖", "3.9100", "0.029082", "15.3975", "0.379187", "干端点来自 ITU 砖；湿端点为本地工程端点。"),
        ("木材", "1.9900", "0.017998", "24.1198", "0.787949", "干端点来自 ITU 木材；湿端点为本地工程端点。"),
        ("玻璃", "6.3100", "0.019276", "6.3100", "0.019276", "非多孔材料，第一版不做内部湿润修正。"),
        ("金属", "1.0000", "1.0e7", "1.0000", "1.0e7", "非多孔材料，第一版不做内部湿润修正。"),
        ("道路/屋顶沥青", "4.8300", "0.062149", "6.6936", "0.115530", "屋顶沥青以沥青混凝土近似。"),
        ("草地/土壤", "13.2338", "0.269711", "50.9439", "1.583564", "以典型中等干燥地面作为资产槽代理。"),
    ])

    doc.add_heading("4. 第二步：水、冰、雪层怎样加入", level=1)
    doc.add_paragraph("水、冰、雪不是材料内部的一部分，而是表面覆盖层。XML 只能写一组 εr、σ，所以第一版使用一个有界权重，把覆盖层折算成该槽的一组等效参数。")
    equation(doc, "qL = role_factor × [1 − exp(−2π × √εr_L(T) × dL / λ0)]")
    doc.add_paragraph("这里 L 可以是 water、ice 或 snow；λ0 是 3.5 GHz 的自由空间波长，约 85.7 mm。dL 必须换算成米。qL 在 0 到 1 之间：覆盖层越厚、电学影响越强，但不会无限超过覆盖层本身。")
    equation(doc, "εr_next = (1 − qL) × εr_previous + qL × εr_L(T)")
    equation(doc, "σ_next  = (1 − qL) × σ_previous  + qL × σ_L(T)")
    doc.add_paragraph("按 ice → water → snow 的顺序重复上述两式；没有某层时，该层厚度为 0，qL 也为 0。")

    doc.add_heading("5. 哪些表面允许有水、冰、雪", level=1)
    table(doc, ["role", "内部湿润度 u", "dw/dr", "ds/di", "原因"], [
        ("ground", "允许", "使用 dw", "允许", "水平地面可以积水、结冰和积雪。"),
        ("roof", "允许", "使用 dr", "允许但可按屋顶保留系数减小", "屋顶会受潮，但坡度会排水、减少覆盖。"),
        ("wall", "多孔墙体允许", "0", "0", "竖直面可能受潮，但第一版不把它当作有毫米级积水或积雪。"),
        ("glass / metal", "通常不变", "0", "0", "非多孔竖直面默认没有内部吸水，也不积水。"),
    ])

    doc.add_heading("6. 材质信息从哪里来", level=1)
    doc.add_paragraph("BigCity 当前槽名已经提供了大部分线索，不是完全无信息。例如 `Roof_Bitumen` 归为屋顶沥青，`Roof_Asphalt` 归为屋顶沥青，`Asphalt1_Road` 归为道路沥青，`Grass1_Ground` 归为草地/土壤。带 `itu_concrete`、`itu_brick`、`itu_glass`、`itu_metal`、`itu_wood` 的槽直接使用对应 ITU 材料类别。")
    note(doc, "仍需诚实保留的假设", "Roof_Bitumen 使用沥青混凝土作为近似；Grass1_Ground 使用典型中等干燥地面作为近似。它们是有来源名称支持的工程代理，不是对该资产真实配方的测量结果。以后有现场材料信息或测量值时，只替换该材料表的端点即可，两个函数不变。")

    doc.add_heading("7. 示例：刚下完大雨，地面积水 1 mm", level=1)
    doc.add_paragraph("输入：T=20°C，u=0.85，dw=1 mm，dr=0 mm，ds=0 mm，di=0 mm。")
    doc.add_paragraph("程序先用第 3 节计算所有槽的 εr_bulk、σ_bulk；再只对 ground 槽施加 1 mm 水层。wall 槽的 dw=0，因此不会得到“墙面也积了 1 mm 水”的错误结果。")
    table(doc, ["材质槽角色", "处理结果"], [
        ("道路沥青 / 地面混凝土 / 草地", "内部湿润度为 0.85，并加入 1 mm 水层。"),
        ("屋顶沥青", "内部湿润度为 0.85；由于 dr=0，不加入屋顶水膜。"),
        ("混凝土外墙 / 砖墙 / 木墙", "内部湿润度为 0.85；不加入积水、冰层或雪层。"),
        ("玻璃 / 金属", "保持其基础参数；不加入内部湿润修正或地面积水。"),
    ])

    doc.add_heading("8. 从对话到 XML 的完整流程", level=1)
    table(doc, ["步骤", "发生什么"], [
        ("1", "用户输入天气描述，例如“大雨刚停，地面积水约 1 mm，20 度”。"),
        ("2", "DeepSeek 只提取 T、u、dw、dr、ds、di 等变量，不输出 εr、σ。"),
        ("3", "本地程序识别每个 XML 材质槽属于哪种 material 和 role。"),
        ("4", "本地程序调用 εr_effective 和 σ_effective，计算全部材质槽。"),
        ("5", "程序生成一份新的动态 XML 和一份 JSON 计算清单。"),
        ("6", "Agent 返回 XML 路径；GUI 将 Sionna 场景切换到该 XML。"),
    ])

    doc.add_heading("9. 参考来源", level=1)
    doc.add_paragraph("1. ITU-R P.2040-4 (2025), Effects of building materials and structures on radio-wave propagation in the range of 1 MHz to 450 GHz：建筑材料干态 εr、σ 的频率模型。")
    doc.add_paragraph("2. ITU-R P.527-6 (2021), Electrical characteristics of the surface of the Earth：液态水、冰、雪随温度变化的复介电参数模型。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
